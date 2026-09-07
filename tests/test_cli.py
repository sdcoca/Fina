"""Tests for fina.cli (WP-9): R-11.3, R-11.4."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl.worksheet.worksheet import Worksheet

from builders import BANK_XLSX, BROKER_CSV, bank_xlsx_with
from fina import cli
from fina.errors import ValidationError
from fina.money import round_half_up
from fina.pipeline import run_pipeline


def _copy_both_fixtures(tmp_path: Path) -> Path:
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    shutil.copyfile(BANK_XLSX, input_dir / "banco_ejemplo.xlsx")
    shutil.copyfile(BROKER_CSV, input_dir / "broker_ejemplo.csv")
    return input_dir


# ---------------------------------------------------------------------------
# T-504 / R-11.3/R-1.24: warnings printed before any figure
# ---------------------------------------------------------------------------


def test_t504_warnings_printed_before_any_figure(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    input_dir = _copy_both_fixtures(tmp_path)
    out_dir = tmp_path / "out"

    exit_code = cli.main(["build", "--input", str(input_dir), "--out", str(out_dir)])

    assert exit_code == 0
    out_lines = capsys.readouterr().out.splitlines()
    warning_indices = [i for i, line in enumerate(out_lines) if line.startswith("warning:")]
    figure_indices = [
        i
        for i, line in enumerate(out_lines)
        if line.startswith(("real_net_worth", "savings_only", "gap"))
    ]
    assert warning_indices, "expected at least one warning in this combined-fixture run"
    assert figure_indices, "expected at least one figure line"
    assert max(warning_indices) < min(figure_indices)


def test_empty_input_directory_succeeds_with_no_figures_printed(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """An empty run (no input files at all) is not an error (R-1.18's spirit extended to the
    whole pipeline): zero entries, an empty Section 1 series, and therefore no figure lines
    at all -- only the `if result.series:` guard's False branch reaches this.
    """
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    out_dir = tmp_path / "out"

    exit_code = cli.main(["build", "--input", str(input_dir), "--out", str(out_dir)])

    assert exit_code == 0
    out = capsys.readouterr().out
    assert "real_net_worth" not in out
    assert (out_dir / "manifest.json").exists()


def test_main_with_no_subcommand_exits_2_with_the_expected_usage_text(
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Pins `argparse.ArgumentParser(prog="fina")` and `add_subparsers(dest="command",
    required=True)` together: the exact program name and the exact `dest` both surface in
    argparse's own generated error text when no subcommand is given.
    """
    with pytest.raises(SystemExit) as exc_info:
        cli.main([])
    assert exc_info.value.code == 2
    err = capsys.readouterr().err
    assert "usage: fina " in err
    assert "fina: error: the following arguments are required: command" in err


def test_build_subcommand_help_text_is_exact(capsys: pytest.CaptureFixture[str]) -> None:
    """Pins the subcommand's one-line `help=` text via an exact stripped-line match, not a
    substring `in` check: a substring check can't tell `"Run the full pipeline..."` apart
    from the same text wrapped in extra characters (e.g. mutmut's own `"XX...XX"` padding),
    since the original text remains a substring of the padded one either way.
    """
    with pytest.raises(SystemExit) as exc_info:
        cli.main(["--help"])
    assert exc_info.value.code == 0
    out = capsys.readouterr().out
    lines = [line.strip() for line in out.splitlines()]
    assert "build     Run the full pipeline over an input directory." in lines


def test_build_without_input_argument_exits_2(tmp_path: Path) -> None:
    """`--input` is `required=True`: omitting it must be an argparse-level usage error (exit
    2), never silently proceed with `input_dir=None`.
    """
    with pytest.raises(SystemExit) as exc_info:
        cli.main(["build", "--out", str(tmp_path / "out")])
    assert exc_info.value.code == 2


def test_build_without_out_argument_exits_2(tmp_path: Path) -> None:
    with pytest.raises(SystemExit) as exc_info:
        cli.main(["build", "--input", str(tmp_path)])
    assert exc_info.value.code == 2


def test_cli_prints_the_final_periods_figures_exactly(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """Ties the printed figures to `result.series[-1]` specifically -- not `[0]`, `[-2]`, or
    a slice off-by-one -- by independently recomputing the expected values via
    `run_pipeline` and `round_half_up`, the same functions `cli._build` itself uses.
    """
    input_dir = _copy_both_fixtures(tmp_path)
    out_dir = tmp_path / "out"

    exit_code = cli.main(["build", "--input", str(input_dir), "--out", str(out_dir)])
    assert exit_code == 0
    out = capsys.readouterr().out

    result = run_pipeline(input_dir)
    latest = result.series[-1]
    assert latest.is_partial  # this fixture combination's final period is genuinely partial
    expected_note = " (partial month)"
    assert (
        f"real_net_worth ({latest.completeness}) as of {latest.as_of.isoformat()}"
        f"{expected_note}: {round_half_up(latest.real_net_worth)} EUR" in out
    )
    assert f"savings_only: {round_half_up(latest.savings_only)} EUR" in out
    assert f"gap: {round_half_up(latest.gap)} EUR" in out


def test_successful_build_writes_manifest_and_chart(tmp_path: Path) -> None:
    input_dir = _copy_both_fixtures(tmp_path)
    out_dir = tmp_path / "out"
    exit_code = cli.main(["build", "--input", str(input_dir), "--out", str(out_dir)])
    assert exit_code == 0
    assert (out_dir / "manifest.json").exists()
    assert (out_dir / "section1_chart.html").exists()


def test_chart_file_contains_the_real_chart_not_the_empty_placeholder(tmp_path: Path) -> None:
    """`render_section1_chart` is called with the *real* prepared rows, not `None`/an empty
    list -- both of which fall into the "no data" placeholder branch and would otherwise
    look like success (a file is still written) while silently containing no chart at all.
    """
    input_dir = _copy_both_fixtures(tmp_path)
    out_dir = tmp_path / "out"
    cli.main(["build", "--input", str(input_dir), "--out", str(out_dir)])
    chart_text = (out_dir / "section1_chart.html").read_text(encoding="utf-8")
    assert "<svg" in chart_text
    assert "No hay datos" not in chart_text


def test_chart_file_is_written_as_utf8_not_the_platform_default(tmp_path: Path) -> None:
    """The chart's Spanish text and euro signs are non-ASCII: writing with the platform's
    default encoding (dropping `encoding="utf-8"` entirely) risks a different byte sequence
    -- or a `UnicodeEncodeError` outright -- on a non-UTF-8 locale.
    """
    input_dir = _copy_both_fixtures(tmp_path)
    out_dir = tmp_path / "out"
    cli.main(["build", "--input", str(input_dir), "--out", str(out_dir)])
    raw = (out_dir / "section1_chart.html").read_bytes()
    assert "€".encode() in raw


def test_cli_prints_correctly_for_a_non_partial_final_month(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """The `else` branch of the partial-month note (an empty string, not a literal
    placeholder) is only reached when the final month is genuinely *not* partial -- the
    combined-fixture tests above never exercise it, since their final month always is.
    """

    def mutate(ws: Worksheet) -> None:
        ws["A9"], ws["B9"] = "31/03/2027", "31/03/2027"  # last calendar day of the month

    input_dir = tmp_path / "input"
    input_dir.mkdir()
    bank_xlsx_with(input_dir, mutate, filename="banco_full_month.xlsx")
    out_dir = tmp_path / "out"

    exit_code = cli.main(["build", "--input", str(input_dir), "--out", str(out_dir)])
    assert exit_code == 0
    out = capsys.readouterr().out

    result = run_pipeline(input_dir)
    latest = result.series[-1]
    assert not latest.is_partial
    assert (
        f"real_net_worth ({latest.completeness}) as of {latest.as_of.isoformat()}: "
        f"{round_half_up(latest.real_net_worth)} EUR" in out
    )


# ---------------------------------------------------------------------------
# T-505 / R-11.4: ReconciliationError aborts with non-zero exit, no report written
# ---------------------------------------------------------------------------


def test_t505_reconciliation_error_aborts_with_no_report_written(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    def mutate(ws: Worksheet) -> None:
        ws["E10"] = "6.196,16€"  # was 6.196,15€: declared balance off by exactly one cent

    input_dir = tmp_path / "input"
    input_dir.mkdir()
    bank_xlsx_with(input_dir, mutate, filename="banco_corrupted.xlsx")
    out_dir = tmp_path / "out"

    exit_code = cli.main(["build", "--input", str(input_dir), "--out", str(out_dir)])

    assert exit_code == 1
    assert not out_dir.exists() or list(out_dir.iterdir()) == []
    assert "error:" in capsys.readouterr().err


# ---------------------------------------------------------------------------
# T-506 / R-11.4: ValidationError aborts the same way
# ---------------------------------------------------------------------------


def test_t506_validation_error_aborts_with_no_report_written(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    """The business logic that *raises* ValidationError (R-3.5's zero-cash-effect internal
    transfer) is already exhaustively tested at its source in test_classification.py; this
    test's job is only to confirm the CLI propagates it the same way it propagates
    ReconciliationError (T-505), which is why the pipeline call is substituted directly.
    """

    def _raise_validation_error(_input_dir: Path, _out_dir: Path | None = None) -> None:
        raise ValidationError(source_file="x.xlsx", source_row=9, invariant="R-3.5: test")

    monkeypatch.setattr(cli, "run_pipeline", _raise_validation_error)

    input_dir = tmp_path / "input"
    input_dir.mkdir()
    out_dir = tmp_path / "out"

    exit_code = cli.main(["build", "--input", str(input_dir), "--out", str(out_dir)])

    assert exit_code == 1
    assert not out_dir.exists() or list(out_dir.iterdir()) == []
    assert "error:" in capsys.readouterr().err
